"""PPTX 解析（v2）：python-pptx 逐页提取文本、表格、备注页、图片。

P2 增强：
- 备注页提取：提取演讲者备注内容
- 内嵌图片：提取幻灯片中的图片，调用 VLM 生成描述
"""
from pptx import Presentation
from langchain_core.documents import Document
from core.vlm import get_vlm_client


def parse_pptx(file_path: str) -> list[Document]:
    """解析 PPTX 为 Markdown 文档块。

    每页幻灯片作为一个 Document，包含：
    - 标题（如果有）
    - 文本内容（按形状顺序）
    - 表格（转为 Markdown 格式）
    - 备注页（演讲者备注）
    - 图片描述（调用 VLM）

    Args:
        file_path: PPTX 文件路径

    Returns:
        Document 列表，每个 Document 代表一页幻灯片
    """
    prs = Presentation(file_path)
    vlm = get_vlm_client()
    docs = []

    for slide_num, slide in enumerate(prs.slides, 1):
        content_parts = []
        image_count = 0

        # 提取标题
        if slide.shapes.title:
            title = slide.shapes.title.text.strip()
            if title:
                content_parts.append(f"# {title}\n")

        # 遍历所有形状
        for shape in slide.shapes:
            # 文本框
            if shape.has_text_frame:
                if shape != slide.shapes.title:  # 跳过标题（已处理）
                    text = shape.text_frame.text.strip()
                    if text:
                        content_parts.append(text)

            # 表格
            elif shape.has_table:
                table_md = _convert_table_to_markdown(shape.table)
                if table_md:
                    content_parts.append(table_md)

            # 图片
            elif shape.shape_type == 13:  # MSO_SHAPE_TYPE.PICTURE
                if vlm.enabled:
                    try:
                        image = shape.image
                        image_bytes = image.blob
                        image_count += 1

                        # 调用 VLM 生成图片描述
                        description = vlm.process_inline_image_sync(image_bytes)
                        if description:
                            content_parts.append(f"[图片{image_count}: {description}]")
                    except Exception:
                        content_parts.append(f"[图片{image_count}: 无法提取描述]")

        # 提取备注页
        if slide.has_notes_slide:
            notes_slide = slide.notes_slide
            notes_text = notes_slide.notes_text_frame.text.strip()
            if notes_text:
                content_parts.append(f"\n**备注：**\n{notes_text}")

        # 组合内容
        if content_parts:
            page_content = "\n\n".join(content_parts)
            docs.append(Document(
                page_content=page_content,
                metadata={"page": slide_num}
            ))

    return docs


def _convert_table_to_markdown(table) -> str:
    """将 PPTX 表格转换为 Markdown 格式。

    Args:
        table: pptx 表格对象

    Returns:
        Markdown 格式的表格字符串
    """
    if not table.rows:
        return ""

    markdown_lines = []

    # 表头
    header_row = table.rows[0]
    header_cells = [cell.text.strip() for cell in header_row.cells]
    markdown_lines.append("| " + " | ".join(header_cells) + " |")

    # 分隔行
    markdown_lines.append("| " + " | ".join(["---"] * len(header_cells)) + " |")

    # 数据行
    for row in table.rows[1:]:
        cells = [cell.text.strip() for cell in row.cells]
        markdown_lines.append("| " + " | ".join(cells) + " |")

    return "\n".join(markdown_lines)
