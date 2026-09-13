import csv

import pymupdf
from docx import Document as DocxDocument
from openpyxl import Workbook
from pptx import Presentation

from core.parsers import load_documents


def test_csv_is_normalized_to_gfm_with_original_columns(tmp_path):
    path = tmp_path / "data.csv"
    with path.open("w", encoding="utf-8-sig", newline="") as f:
        writer = csv.writer(f)
        writer.writerow(["城市", "人数"])
        writer.writerow(["北京", "10"])
        writer.writerow(["上海", "20"])

    docs = load_documents(str(path))
    text = "\n".join(d.page_content for d in docs)

    assert "| 城市 | 人数 |" in text
    assert "| 北京 | 10 |" in text
    assert "| 上海 | 20 |" in text


def test_xlsx_keeps_sheet_name_and_columns(tmp_path):
    path = tmp_path / "data.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.title = "指标"
    ws.append(["项目", "数值"])
    ws.append(["召回率", 0.9])
    wb.save(path)

    docs = load_documents(str(path))
    text = "\n".join(d.page_content for d in docs)

    assert "# Sheet: 指标" in text
    assert "| 项目 | 数值 |" in text
    assert "| 召回率 | 0.9 |" in text


def test_html_preserves_heading_list_and_table(tmp_path):
    path = tmp_path / "page.html"
    path.write_text(
        """
        <html><body>
          <h1>系统说明</h1>
          <p>第一段。</p>
          <ul><li>列表一</li><li>列表二</li></ul>
          <table>
            <tr><th>字段</th><th>含义</th></tr>
            <tr><td>name</td><td>名称</td></tr>
          </table>
        </body></html>
        """,
        encoding="utf-8",
    )

    text = load_documents(str(path))[0].page_content

    assert "# 系统说明" in text
    assert "- 列表一" in text
    assert "| 字段 | 含义 |" in text
    assert "| name | 名称 |" in text


def test_docx_markdown_and_pptx_keep_structure(tmp_path):
    docx_path = tmp_path / "sample.docx"
    doc = DocxDocument()
    doc.add_heading("标题", level=1)
    doc.add_paragraph("正文")
    table = doc.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "字段"
    table.cell(0, 1).text = "值"
    table.cell(1, 0).text = "A"
    table.cell(1, 1).text = "1"
    doc.save(docx_path)

    docx_text = "\n".join(d.page_content for d in load_documents(str(docx_path)))
    assert "# 标题" in docx_text
    assert "| 字段 | 值 |" in docx_text

    pptx_path = tmp_path / "sample.pptx"
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[5])
    slide.shapes.title.text = "幻灯片标题"
    prs.save(pptx_path)

    pptx_docs = load_documents(str(pptx_path))
    assert pptx_docs[0].metadata["page"] == 1
    assert "# 幻灯片标题" in pptx_docs[0].page_content


def test_pdf_keeps_page_number(tmp_path):
    path = tmp_path / "sample.pdf"
    pdf = pymupdf.open()
    page = pdf.new_page()
    page.insert_text((72, 72), "PDF page content")
    pdf.save(path)
    pdf.close()

    docs = load_documents(str(path))

    assert docs
    assert docs[0].metadata["page"] == 1
