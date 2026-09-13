from pathlib import Path

from docx import Document as DocxDocument
from PIL import Image

from core.parsers.docx_parser import _extract_images_from_paragraph


def _paragraph_images(paragraph, doc) -> list[bytes]:
    return _extract_images_from_paragraph(paragraph, doc)


def test_docx_image_extraction_only_returns_images_from_current_paragraph(tmp_path: Path):
    image_path = tmp_path / "pixel.png"
    Image.new("RGB", (2, 2), "red").save(image_path)

    doc = DocxDocument()
    first = doc.add_paragraph("first")
    first.add_run().add_picture(str(image_path))
    second = doc.add_paragraph("second")
    third = doc.add_paragraph("third")
    third.add_run().add_picture(str(image_path))

    assert len(_paragraph_images(first, doc)) == 1
    assert _paragraph_images(second, doc) == []
    assert len(_paragraph_images(third, doc)) == 1
